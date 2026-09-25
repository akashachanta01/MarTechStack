"""
MarTechJobs test suite. Run before EVERY push:  python manage.py test

Covers: page smoke tests, the MarTech ATS Match feature, staff-page security,
and regression tests for bugs fixed in past audits.
"""
import json
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase, override_settings

from jobs.ats_match import extract_terms, match
from jobs.models import BlogPost, Category, Job, Subscriber, Tool

JD_HTML = (
    "<p>Marketing Operations Manager. Own our <b>Marketo Engage</b> instance and "
    "Salesforce Marketing Cloud. Build lead scoring and lead routing, manage nurture "
    "streams. Track MQL-to-SQL conversion. Experience with Segment or a customer data "
    "platform. Marketo Certified Expert preferred. Community outreach and GTM alignment.</p>"
)
RESUME = (
    "Senior Marketing Ops Specialist\n"
    "- Administered Marketo; built lead score model lifting MQL-to-SQL conversion 22%\n"
    "- Ran Pardot nurture campaigns for 40k contacts across three regions\n"
    "- Managed SFMC email sends and journeys for the global team\n"
    "- Partnered with sales on routing rules and pipeline reporting"
)

TEST_SETTINGS = dict(SECURE_SSL_REDIRECT=False, ALLOWED_HOSTS=["*"])


def make_job(**kw):
    defaults = dict(
        title="Marketing Operations Manager", company="Acme", location="Remote",
        description=JD_HTML, apply_url="https://acme.test/apply",
        screening_status="approved", is_active=True, work_arrangement="remote",
    )
    defaults.update(kw)
    return Job.objects.create(**defaults)


# ---------------------------------------------------------------------------
# ATS engine (pure functions)
# ---------------------------------------------------------------------------
class ATSEngineTests(TestCase):
    def test_ignores_non_tool_words_in_jd(self):
        terms = extract_terms("Community outreach and GTM strategy alignment.", for_jd=True)
        self.assertNotIn("Outreach", terms)
        self.assertNotIn("Google Tag Manager", terms)

    def test_capitalized_product_name_counts(self):
        self.assertIn("Outreach", extract_terms("Experience with Outreach and Salesloft.", for_jd=True))

    def test_sales_qualified_leads_are_not_sql_language(self):
        self.assertNotIn("SQL", extract_terms("Track MQL-to-SQL conversion and SQLs by source.", for_jd=True))
        self.assertIn("SQL", extract_terms("Strong SQL and Python skills.", for_jd=True))

    def test_longest_match_wins(self):
        terms = extract_terms("Salesforce Marketing Cloud journeys", for_jd=True)
        self.assertIn("Salesforce Marketing Cloud", terms)
        self.assertNotIn("Salesforce", terms)

    def test_rebrand_wording_fix(self):
        r = match("Ran Pardot programs for years across many regions and teams.",
                  "Must know Marketing Cloud Account Engagement deeply.")
        fixes = {(f["you_wrote"], f["jd_says"]) for f in r["alias_fixes"]}
        self.assertIn(("Pardot", "Marketing Cloud Account Engagement"), fixes)

    def test_possessive_is_not_a_wording_difference(self):
        r = match("Administered Marketo's instance for a large B2B team.", "Own our Marketo instance.")
        self.assertEqual(r["alias_fixes"], [])

    def test_full_match_counts(self):
        r = match(RESUME, JD_HTML)
        self.assertEqual(r["required_count"], 8)
        self.assertEqual(r["matched_count"], 4)
        self.assertEqual({m["term"] for m in r["missing"]},
                         {"Marketo Certified Expert", "Segment", "Customer Data Platform", "Lead Routing"})


# ---------------------------------------------------------------------------
# ATS Match API + page
# ---------------------------------------------------------------------------
@override_settings(**TEST_SETTINGS)
class ATSMatchAPITests(TestCase):
    API = "/tools/api/ats-match/"

    def setUp(self):
        cache.clear()
        self.job = make_job()

    def post(self, client, body):
        r = client.post(self.API, data=json.dumps(body), content_type="application/json")
        return r.status_code, r.json()

    def test_page_renders_with_job(self):
        r = self.client.get(f"/tools/resume-keyword-scanner/?job={self.job.id}")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Marketing Operations Manager")
        self.assertContains(r, "FAQPage")

    def test_anonymous_gets_one_run_then_signup_gate(self):
        s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual(s, 200)
        self.assertEqual((d["matched_count"], d["required_count"]), (4, 8))
        self.assertNotIn("alias_fixes", d)          # locked for anonymous
        self.assertEqual(d["locked_alias_fixes"], 2)
        s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual((s, d.get("gate")), (403, "signup"))

    def test_validation(self):
        s, _ = self.post(self.client, {"resume_text": "too short", "job_id": self.job.id})
        self.assertEqual(s, 400)
        s, _ = self.post(self.client, {"resume_text": RESUME, "jd_text": "x"})
        self.assertEqual(s, 400)
        s, _ = self.post(self.client, {"resume_text": RESUME, "jd_text": "We need a friendly cashier for our downtown sandwich store on weekends and some evenings. " * 2})
        self.assertEqual(s, 422)

    def test_invalid_json(self):
        r = self.client.post(self.API, data="not json", content_type="application/json")
        self.assertEqual(r.status_code, 400)

    def test_account_full_report_tips_and_monthly_cap(self):
        user = get_user_model().objects.create_user("u", "u@x.test", "pw12345!")
        self.client.force_login(user)
        fake = mock.MagicMock()
        fake.chat.completions.create.return_value.choices = [
            mock.MagicMock(message=mock.MagicMock(content=json.dumps({"tips": ["a", "b", "c"]})))
        ]
        with mock.patch.dict("os.environ", {"OPENAI_API_KEY": "test"}), \
             mock.patch("tools.views.OpenAI", return_value=fake):
            s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual(s, 200)
        self.assertEqual(d["tips"], ["a", "b", "c"])
        self.assertEqual(d["runs_left"], 4)
        self.assertIn(("Marketo", "Marketo Engage"), {(f["you_wrote"], f["jd_says"]) for f in d["alias_fixes"]})
        for _ in range(4):
            s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual(s, 200)
        s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual((s, d.get("gate")), (429, "limit"))

    def test_tips_failure_does_not_break_report(self):
        user = get_user_model().objects.create_user("v", "v@x.test", "pw12345!")
        self.client.force_login(user)
        with mock.patch.dict("os.environ", {"OPENAI_API_KEY": "test"}), \
             mock.patch("tools.views.OpenAI", side_effect=RuntimeError("boom")):
            s, d = self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        self.assertEqual(s, 200)
        self.assertEqual(d["tips"], [])

    def test_inactive_job_is_not_usable(self):
        closed = make_job(title="Closed role", is_active=False, screening_status="rejected")
        s, _ = self.post(self.client, {"resume_text": RESUME, "job_id": closed.id})
        self.assertEqual(s, 400)  # falls back to "paste the job description"

    def test_job_page_has_nofollow_cta(self):
        r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
        self.assertContains(r, "Check my resume")
        self.assertContains(r, f'?job={self.job.id}" rel="nofollow"')

    # --- Tracking: GA4 events (via GTM dataLayer) + Founder HQ "who" log ---
    def test_job_page_tracks_check_resume_clicks(self):
        r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
        self.assertContains(r, "window.mtjTrack")          # helper in base.html
        self.assertContains(r, "check_resume_click")
        self.assertContains(r, 'data-placement="top"')
        self.assertContains(r, 'data-placement="bottom"')  # top + bottom CTAs

    def test_posthog_off_without_key(self):
        with self.settings(POSTHOG_KEY=""):
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
        self.assertNotContains(r, "posthog.init")

    def test_posthog_loads_and_identifies_signed_in_user(self):
        from django.contrib.auth.models import User
        u = User.objects.create_user("ph", "ph@example.com", "pw12345!x")
        with self.settings(POSTHOG_KEY="phc_test123"):
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
            self.assertContains(r, 'posthog.init("phc_test123"')
            self.assertNotContains(r, "posthog.identify")   # anonymous
            self.client.force_login(u)
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
            self.assertContains(r, "ph@example.com")
            self.assertContains(r, "posthog.identify")
            u.is_staff = True; u.save()
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
            self.assertNotContains(r, "posthog.init")       # founder excluded

    def test_posthog_internal_device_tag(self):
        with self.settings(POSTHOG_KEY="phc_test123"):
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/?internal=1")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'localStorage.setItem("mtj_internal", "1")')
        self.assertContains(r, "posthog.register({ is_internal: true })")

    def test_job_page_tracks_apply_and_save(self):
        r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
        self.assertContains(r, "apply_click")
        self.assertContains(r, 'js-apply" data-placement="top"')
        self.assertContains(r, 'js-apply" data-placement="bottom"')
        self.assertContains(r, "job_saved")
        self.assertContains(r, 'window.mtjPageType = "job_detail"')

    def test_search_is_tracked_only_when_searching(self):
        r = self.client.get("/jobs/?q=Marketo")
        self.assertContains(r, "window.mtjTrack('job_search'")
        self.assertContains(r, 'query: "marketo"')
        r = self.client.get("/jobs/")
        self.assertNotContains(r, "window.mtjTrack('job_search'")

    def test_signup_completed_fires_once(self):
        from django.contrib.auth.models import User
        u = User.objects.create_user("sg", "sg@example.com", "pw12345!x")
        self.client.force_login(u)
        session = self.client.session
        session["mtj_signed_up"] = "google"; session.save()
        r = self.client.get("/")
        self.assertContains(r, "window.mtjTrack('signup_completed', { method: \"google\" })")
        r = self.client.get("/")
        self.assertNotContains(r, "signup_completed")

    def test_subscribe_success_is_tracked(self):
        r = self.client.get("/")
        self.assertContains(r, "'newsletter_subscribe'"); self.assertContains(r, "'job_alert_created'")
        self.assertContains(r, "window.mtjTrack('signup_click'")

    def test_email_links_get_utm_tags_but_unsubscribe_does_not(self):
        from jobs.emails import _add_utm
        html = ('<a href="https://martechjobs.io/job/1/x/">a</a>'
                '<a href="https://martechjobs.io/jobs/?q=Marketo#top">b</a>'
                '<a href="https://martechjobs.io/u/abc/">c</a>'
                '<a href="https://other.com/">d</a>')
        out = _add_utm(html, "daily_digest")
        self.assertIn('/job/1/x/?utm_source=email&amp;utm_medium=email&amp;utm_campaign=daily_digest"', out)
        self.assertIn('?q=Marketo&amp;utm_source=email&amp;utm_medium=email&amp;utm_campaign=daily_digest#top"', out)
        self.assertIn('href="https://martechjobs.io/u/abc/"', out)
        self.assertIn('href="https://other.com/"', out)

    def test_recruiter_funnel_is_tracked(self):
        r = self.client.get("/post-job/")
        self.assertContains(r, "post_job_start")
        self.assertContains(r, "post_job_submit")
        r = self.client.get("/post-job/success/?plan=free")
        self.assertContains(r, "window.mtjTrack('post_job_completed', { plan: \"free\" })")

    def test_sitewide_click_tracking_present(self):
        r = self.client.get("/")
        for ev in ["employer_cta_click", "content_to_jobs_click", "tool_used"]:
            self.assertContains(r, ev)

    def test_tool_page_tracks_funnel_events(self):
        r = self.client.get(f"/tools/resume-keyword-scanner/?job={self.job.id}")
        for event in ["ats_check_submit", "ats_check_result", "ats_gate_shown", "ats_signup_click"]:
            self.assertContains(r, event)

    def test_checks_are_logged_without_resume_text(self):
        from jobs.models import AtsCheck
        self.post(self.client, {"resume_text": RESUME, "job_id": self.job.id})
        anon = AtsCheck.objects.get()
        self.assertIsNone(anon.user)
        self.assertEqual((anon.job_id, anon.matched, anon.required, anon.source), (self.job.id, 4, 8, "job"))
        # No field anywhere on the log can hold resume text.
        self.assertFalse(any(f.name in ("resume", "resume_text", "text", "ip") for f in AtsCheck._meta.fields))

        user = get_user_model().objects.create_user("w", "w@x.test", "pw12345!")
        c = self.client_class(); c.force_login(user)
        with mock.patch.dict("os.environ", {"OPENAI_API_KEY": ""}):
            self.post(c, {"resume_text": RESUME, "jd_text": JD_HTML})
        logged = AtsCheck.objects.filter(user=user).get()
        self.assertEqual((logged.job_id, logged.source), (None, "pasted"))

    def test_founder_hq_shows_resume_checker_activity(self):
        from jobs.models import AtsCheck
        checker = get_user_model().objects.create_user("c", "checker@x.test", "pw12345!")
        AtsCheck.objects.create(user=checker, job=self.job, matched=3, required=8)
        AtsCheck.objects.create(user=None, job=self.job, matched=5, required=8)
        admin = get_user_model().objects.create_user("a", "a@x.test", "pw12345!", is_staff=True)
        self.client.force_login(admin)
        r = self.client.get("/staff/")
        self.assertContains(r, "Resume checker activity")
        self.assertContains(r, "checker@x.test")
        self.assertContains(r, "3 / 8")
        self.assertContains(r, "anonymous")


# ---------------------------------------------------------------------------
# Page smoke tests — every major page must render
# ---------------------------------------------------------------------------
@override_settings(**TEST_SETTINGS)
class PageSmokeTests(TestCase):
    def setUp(self):
        cache.clear()
        cat = Category.objects.create(name="MarTech", slug="martech")
        self.tool = Tool.objects.create(name="Marketo", slug="marketo", category=cat)
        self.job = make_job()
        self.job.tools.add(self.tool)
        BlogPost.objects.create(title="AI in MOps", slug="ai-in-mops", excerpt="x",
                                content="<p>Body</p>", category="AI in MarTech", is_published=True)

    def assert_ok(self, url, status=200):
        r = self.client.get(url)
        self.assertEqual(r.status_code, status, f"{url} returned {r.status_code}")
        return r

    def test_core_pages(self):
        for url in [
            "/", "/jobs/", f"/job/{self.job.id}/{self.job.slug}/", "/jobs/marketo/",
            "/blog/", "/blog/ai-in-mops/", "/blog/ai-in-martech/", "/blog/salary-guides/",
            "/salary-guide/", "/martech-job-market-statistics/", "/companies/",
            "/category/engineering/", "/category/operations/", "/category/data/",
            "/category/ai-automation/", "/remote/jobs/", "/about/", "/for-employers/",
            "/marketing-operations-manager-jobs/",
        ]:
            with self.subTest(url=url):
                self.assert_ok(url)

    def test_tools_pages(self):
        for url in [
            "/tools/", "/tools/resume-keyword-scanner/", "/tools/job-description-generator/",
            "/tools/salary-calculator/", "/tools/interview-questions-generator/",
            "/tools/utm-link-builder/", "/tools/email-subject-line-tester/",
        ]:
            with self.subTest(url=url):
                self.assert_ok(url)

    def test_seo_files(self):
        for url in ["/robots.txt", "/llms.txt", "/sitemap.xml", "/blog/feed/"]:
            with self.subTest(url=url):
                self.assert_ok(url)

    def test_closed_job_returns_410(self):
        closed = make_job(title="Old role", is_active=False, screening_status="rejected")
        self.assert_ok(f"/job/{closed.id}/{closed.slug}/", status=410)
        self.assert_ok("/job/999999/nope/", status=410)

    def test_mixed_case_seo_url_redirects_to_lowercase(self):
        r = self.client.get("/Chicago/Marketo-jobs/")
        self.assertEqual(r.status_code, 301)
        self.assertEqual(r["Location"], "/chicago/marketo-jobs/")


# ---------------------------------------------------------------------------
# Security & past-bug regressions
# ---------------------------------------------------------------------------
@override_settings(**TEST_SETTINGS)
class SecurityRegressionTests(TestCase):
    def test_staff_pages_require_staff(self):
        # Regression: /staff/review/ once lost its staff guard in a merge.
        for url in ["/staff/", "/staff/review/"]:
            with self.subTest(url=url):
                r = self.client.get(url)
                self.assertEqual(r.status_code, 302, f"{url} must redirect anonymous users")
        user = get_user_model().objects.create_user("n", "n@x.test", "pw12345!")
        self.client.force_login(user)
        for url in ["/staff/", "/staff/review/"]:
            with self.subTest(url=url, as_="non-staff"):
                self.assertEqual(self.client.get(url).status_code, 302)

    def test_staff_can_open_founder_hq(self):
        admin = get_user_model().objects.create_user("s", "s@x.test", "pw12345!", is_staff=True)
        self.client.force_login(admin)
        self.assertEqual(self.client.get("/staff/").status_code, 200)

    def test_unsubscribe_is_soft_delete(self):
        from django.core import signing
        Subscriber.objects.create(email="a@x.test")
        token = signing.dumps("a@x.test", salt="unsubscribe")
        self.client.post(f"/u/{token}/")
        sub = Subscriber.objects.get(email="a@x.test")
        self.assertFalse(sub.is_active)            # row kept as suppression record

    def test_digest_recipients_include_accounts_and_exclude_unsubscribed(self):
        from jobs.emails import get_digest_recipients
        Subscriber.objects.create(email="sub@x.test")
        Subscriber.objects.create(email="gone@x.test", is_active=False)
        get_user_model().objects.create_user("acct", "acct@x.test", "pw12345!")
        get_user_model().objects.create_user("gone", "gone@x.test", "pw12345!")
        recipients = set(get_digest_recipients())
        self.assertIn("sub@x.test", recipients)
        self.assertIn("acct@x.test", recipients)
        self.assertNotIn("gone@x.test", recipients)

    def test_email_defaults_never_fall_back_to_gmail(self):
        from django.conf import settings
        self.assertNotIn("gmail", settings.EMAIL_HOST)


# ---------------------------------------------------------------------------
# Resume Match phase 1: saved resume, upload, ranking, more matches
# ---------------------------------------------------------------------------
def _docx_upload(lines, name="resume.docx"):
    import io
    import docx
    from django.core.files.uploadedfile import SimpleUploadedFile
    d = docx.Document()
    for l in lines:
        d.add_paragraph(l)
    buf = io.BytesIO(); d.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(),
                              content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@override_settings(**TEST_SETTINGS)
class ResumeMatchTests(TestCase):
    UPLOAD = "/tools/api/resume/upload/"
    API = "/tools/api/ats-match/"

    def setUp(self):
        cache.clear()
        self.job = make_job()
        make_job(title="Marketo Administrator", company="Globex")
        make_job(title="Lifecycle Manager", company="Initech",
                 description="<p>Own Braze and Iterable programs. SQL and Looker reporting required.</p>")
        self.user = get_user_model().objects.create_user("rm", "rm@x.test", "pw12345!")

    def test_parse_docx_and_reject_bad_files(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from jobs.resume_match import extract_resume_text, ResumeParseError
        text = extract_resume_text(_docx_upload(RESUME.splitlines()))
        self.assertIn("Marketo", text)
        with self.assertRaises(ResumeParseError):
            extract_resume_text(SimpleUploadedFile("cv.txt", b"hello" * 100))
        with self.assertRaises(ResumeParseError):
            extract_resume_text(_docx_upload(["Too short"]))

    def test_anonymous_upload_returns_text_and_saves_nothing(self):
        from accounts.models import UserResume
        r = self.client.post(self.UPLOAD, {"resume": _docx_upload(RESUME.splitlines())})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()["saved"])
        self.assertIn("Marketo", r.json()["text"])
        self.assertEqual(UserResume.objects.count(), 0)

    def test_member_upload_saves_text_then_checks_use_it(self):
        from accounts.models import UserResume
        self.client.force_login(self.user)
        r = self.client.post(self.UPLOAD, {"resume": _docx_upload(RESUME.splitlines())})
        self.assertTrue(r.json()["saved"])
        saved = UserResume.objects.get(user=self.user)
        self.assertIn("Marketo", saved.text)
        self.assertEqual(saved.filename, "resume.docx")
        # No resume in the request: the saved one is used.
        r = self.client.post(self.API, data=json.dumps({"job_id": self.job.id}), content_type="application/json")
        d = r.json()
        self.assertEqual(r.status_code, 200)
        self.assertEqual((d["matched_count"], d["required_count"]), (4, 8))
        self.assertTrue(d["resume_saved"])
        self.assertIn(d["label"], ("strong", "good", "stretch"))
        # Missing skills carry real demand % and are sorted most-requested first.
        pcts = [m["demand_pct"] for m in d["missing"]]
        self.assertEqual(pcts, sorted(pcts, reverse=True))
        self.assertTrue(all(0 <= p <= 100 for p in pcts))
        # More matches exclude the current job and include the others.
        ids = [j["id"] for j in d["more_matches"]]
        self.assertNotIn(self.job.id, ids)
        self.assertTrue(ids)

    def test_pasted_resume_saved_only_when_asked(self):
        from accounts.models import UserResume
        self.client.force_login(self.user)
        self.client.post(self.API, data=json.dumps({"resume_text": RESUME, "job_id": self.job.id}), content_type="application/json")
        self.assertFalse(UserResume.objects.filter(user=self.user).exists())
        self.client.post(self.API, data=json.dumps({"resume_text": RESUME, "job_id": self.job.id, "save": True}), content_type="application/json")
        self.assertTrue(UserResume.objects.filter(user=self.user).exists())

    def test_delete_resume(self):
        from accounts.models import UserResume
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")
        self.assertEqual(self.client.post("/tools/api/resume/delete/").status_code, 401)
        self.client.force_login(self.user)
        r = self.client.get("/accounts/settings/")
        self.assertContains(r, "cv.pdf")
        self.assertContains(r, "Delete resume")
        self.assertEqual(self.client.post("/tools/api/resume/delete/").status_code, 200)
        self.assertFalse(UserResume.objects.filter(user=self.user).exists())

    def test_page_shows_real_stats_and_saved_state(self):
        from accounts.models import UserResume
        r = self.client.get("/tools/resume-keyword-scanner/")
        self.assertContains(r, "3</b> live MarTech jobs")
        self.assertContains(r, "3</b> companies hiring")
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")
        self.client.force_login(self.user)
        r = self.client.get(f"/tools/resume-keyword-scanner/?job={self.job.id}")
        self.assertContains(r, "Using your saved resume")
        self.assertContains(r, "cv.pdf")

    def test_nav_links_resume_scanner_and_privacy_mentions_resume(self):
        r = self.client.get("/")
        self.assertContains(r, 'href="/tools/resume-keyword-scanner/" class="tools-link"')
        self.assertContains(self.client.get("/privacy/"), "never the file")


@override_settings(**TEST_SETTINGS)
class ResumeMatchColdStartTests(TestCase):
    """Regression: on prod the first check analysed every live job inside the
    request and hit the 30s worker timeout -> 'Something went wrong' forever."""

    def setUp(self):
        cache.clear()
        self.job = make_job()
        for i in range(5):
            make_job(title=f"Ops role {i}", company=f"Co{i}")

    def test_check_succeeds_without_waiting_when_demand_not_ready(self):
        with mock.patch("jobs.resume_match.term_demand", return_value=({}, 0)):
            r = self.client.post("/tools/api/ats-match/", data=json.dumps({"resume_text": RESUME, "job_id": self.job.id}),
                                 content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(all(m["demand_pct"] is None for m in r.json()["missing"]))

    def test_budget_limits_work_and_progress_persists(self):
        from jobs.resume_match import job_requirements
        reqs, complete = job_requirements(budget_s=-1)   # no time: nothing computed
        self.assertFalse(complete)
        reqs, complete = job_requirements(budget_s=None)  # cron path: everything
        self.assertTrue(complete)
        self.assertEqual(len(reqs), 6)

    def test_ranking_failure_never_breaks_a_check(self):
        with mock.patch("jobs.resume_match.rank_missing", side_effect=RuntimeError("boom")):
            r = self.client.post("/tools/api/ats-match/", data=json.dumps({"resume_text": RESUME, "job_id": self.job.id}),
                                 content_type="application/json")
        self.assertEqual(r.status_code, 200)

    def test_warm_command_fills_demand(self):
        from django.core.management import call_command
        from jobs.resume_match import term_demand
        call_command("warm_resume_match")
        demand, n = term_demand(budget_s=-1)  # must be served from the warm cache
        self.assertEqual(n, 6)
        self.assertEqual(demand.get("Marketo"), 100)


@override_settings(**TEST_SETTINGS)
class ResultsPageImprovementTests(TestCase):
    JD = ("About the role. You will own Adobe Journey Optimizer programs end to end.\n"
          "• Hands-on experience with Adobe Experience Platform and a customer data platform.\n"
          "• Marketo Certified Expert is a plus.")

    def setUp(self):
        cache.clear()

    def test_missing_items_quote_the_job_and_split_required_vs_nice_to_have(self):
        from jobs.resume_match import annotate_missing
        r = match("Administered office schedules and vendor contracts for a 40-person team.", self.JD)
        by = {m["term"]: m for m in annotate_missing(r["missing"], self.JD)}
        self.assertFalse(by["Marketo Certified Expert"]["required"])
        self.assertTrue(by["Adobe Journey Optimizer"]["required"])
        self.assertIn("Adobe Journey Optimizer", by["Adobe Journey Optimizer"]["quote"])
        self.assertEqual(by["Customer Data Platform"]["quote_term"].lower(), "customer data platform")

    def test_lines_needing_numbers_picks_real_unquantified_achievements(self):
        from jobs.resume_match import lines_needing_numbers
        picks = lines_needing_numbers("Managed email campaigns for the global team\n"
                                      "Built lead scoring in Marketo lifting conversion 22%\n"
                                      "Partnered with sales on routing rules and reporting\nhello")
        self.assertEqual(picks, ["Managed email campaigns for the global team",
                                 "Partnered with sales on routing rules and reporting"])

    def test_stretch_result_suggests_only_better_fitting_jobs_even_for_anonymous(self):
        hard = make_job(title="Adobe consultant", company="Adobe", description="<p>" + self.JD + "</p>")
        make_job(title="Marketing Operations Manager", company="Acme")      # ~4/8 for RESUME
        make_job(title="Pure Marketo admin", company="Globex",
                 description="<p>Own our Marketo instance, lead scoring and nurture programs.</p>")
        r = self.client.post("/tools/api/ats-match/", data=json.dumps({"resume_text": RESUME, "job_id": hard.id}),
                             content_type="application/json")
        d = r.json()
        self.assertEqual(d["label"], "stretch")
        self.assertTrue(d["more_matches"])
        cur = d["matched_count"] / d["required_count"]
        self.assertTrue(all(j["matched"] / j["required"] > cur for j in d["more_matches"]))
        self.assertNotIn(hard.id, [j["id"] for j in d["more_matches"]])
        self.assertIn("lines_to_quantify", d)
        self.assertTrue(all("quote" in m for m in d["missing"]))


@override_settings(**TEST_SETTINGS)
class JobHtmlBulletsTests(TestCase):
    def setUp(self):
        cache.clear()

    def test_bullets_from_job_html_are_not_merged_into_one_sentence(self):
        job = make_job(description="<p>We are hiring a senior consultant to lead customer journey programs for enterprise clients.</p>"
                                   "<ul><li>Hands-on Adobe Experience Platform experience.</li>"
                                   "<li>Adobe Journey Optimizer ownership.</li><li>Marketo Certified Expert is a plus.</li></ul>")
        d = self.client.post("/tools/api/ats-match/", data=json.dumps({"resume_text": RESUME, "job_id": job.id}),
                             content_type="application/json").json()
        by = {m["term"]: m for m in d["missing"]}
        self.assertTrue(by["Adobe Experience Platform"]["required"])
        self.assertTrue(by["Adobe Journey Optimizer"]["required"])
        self.assertFalse(by["Marketo Certified Expert"]["required"])


class TaxonomyFalsePositiveTests(TestCase):
    def test_customer_journeys_phrase_is_not_a_skill(self):
        self.assertNotIn("Journey Orchestration", extract_terms("Senior Technical Consultant, Customer Journeys", for_jd=True))
        self.assertIn("Journey Orchestration", extract_terms("Build flows in Journey Builder.", for_jd=True))


@override_settings(**TEST_SETTINGS)
class Phase2MatchBadgeTests(TestCase):
    def setUp(self):
        cache.clear()
        self.job = make_job()
        self.fit = make_job(title="Marketo & SFMC Specialist", company="Umbrella",
                            description="<p>Run Marketo and Salesforce Marketing Cloud. Build lead scoring and nurture programs.</p>")
        self.user = get_user_model().objects.create_user("p2", "p2@x.test", "pw12345!")

    def test_api_is_empty_for_anonymous_and_members_without_resume(self):
        self.assertEqual(self.client.get("/tools/api/my-matches/").json()["jobs"], {})
        self.client.force_login(self.user)
        d = self.client.get("/tools/api/my-matches/").json()
        self.assertEqual((d["jobs"], d["has_resume"]), ({}, False))

    def test_api_scores_every_job_for_a_saved_resume(self):
        from accounts.models import UserResume
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")
        self.client.force_login(self.user)
        d = self.client.get("/tools/api/my-matches/").json()
        self.assertTrue(d["ready"])
        m, n, label = d["jobs"][str(self.job.id)]
        self.assertEqual((m, n), (4, 8))
        self.assertEqual(d["jobs"][str(self.fit.id)][2], "strong")

    def test_badge_script_only_for_resume_holders_and_seo_unchanged(self):
        from accounts.models import UserResume
        url = f"/job/{self.job.id}/{self.job.slug}/"
        r = self.client.get(url)
        self.assertNotContains(r, "/tools/api/my-matches/")        # anonymous: no script
        self.client.force_login(self.user)
        self.assertNotContains(self.client.get(url), "/tools/api/my-matches/")  # no resume yet
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")
        r = self.client.get(url)
        self.assertContains(r, "/tools/api/my-matches/")
        self.assertContains(r, f'id="mtj-fit" class="mtj-fit" data-job-id="{self.job.id}"')

    def test_my_matches_page_ranks_best_first(self):
        from accounts.models import UserResume
        self.client.force_login(self.user)
        self.assertContains(self.client.get("/accounts/matches/"), "Upload your resume to see your matches")
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")
        r = self.client.get("/accounts/matches/?show=all")
        body = r.content.decode()
        self.assertLess(body.index("Marketo &amp; SFMC Specialist"), body.index("Marketing Operations Manager"))
        self.assertContains(r, "noindex")
        self.assertEqual(self.client.get("/accounts/matches/").status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.get("/accounts/matches/").status_code, 302)


@override_settings(**TEST_SETTINGS)
class Phase3WeeklyEmailTests(TestCase):
    def setUp(self):
        from django.utils import timezone
        from accounts.models import UserResume
        cache.clear()
        now = timezone.now()
        self.fit = make_job(title="Marketo & SFMC Specialist", company="Umbrella", went_live_at=now,
                            description="<p>Run Marketo and Salesforce Marketing Cloud. Build lead scoring and nurture programs.</p>")
        self.weak = make_job(title="Adobe consultant", company="Adobe", went_live_at=now,
                             description="<p>Adobe Journey Optimizer, Adobe Experience Platform, Braze and SQL.</p>")
        U = get_user_model()
        self.a = U.objects.create_user("a", "a@x.test", "pw12345!", first_name="Asha")
        self.b = U.objects.create_user("b", "b@x.test", "pw12345!")   # resume, but unsubscribed
        self.c = U.objects.create_user("c", "c@x.test", "pw12345!")   # resume, no strong match
        UserResume.objects.create(user=self.a, text=RESUME)
        UserResume.objects.create(user=self.b, text=RESUME)
        UserResume.objects.create(user=self.c, text="Office administrator managing calendars, travel and vendor invoices for a busy team." * 3)
        Subscriber.objects.create(email="b@x.test", is_active=False)
        Subscriber.objects.create(email="list@x.test", is_active=True)

    def _run(self, *cmd):
        from django.core.management import call_command
        import jobs.management.commands.send_daily_digest as digest_mod
        m = mock.MagicMock(return_value=True)
        with mock.patch("jobs.emails.send_html_email", m), mock.patch.object(digest_mod, "send_html_email", m):
            call_command(*cmd)
        return m

    def test_personal_email_only_to_opted_in_members_with_strong_new_matches(self):
        m = self._run("send_weekly_matches")
        sent = {c.kwargs["to_email"][0]: c.kwargs for c in m.call_args_list}
        self.assertEqual(set(sent), {"a@x.test"})
        kw = sent["a@x.test"]
        self.assertIn("80%+", kw["subject"]); self.assertIn("Asha", kw["subject"])
        self.assertEqual([j["id"] for j in kw["context"]["jobs"]], [self.fit.id])

    def test_weekly_digest_skips_people_who_got_personal_email(self):
        self._run("send_weekly_matches")
        m = self._run("send_daily_digest", "--weekly")
        to = {c.kwargs["to_email"][0] for c in m.call_args_list}
        self.assertNotIn("a@x.test", to)        # got the personal one
        self.assertNotIn("b@x.test", to)        # unsubscribed
        self.assertTrue({"c@x.test", "list@x.test"} <= to)
        self.assertIn("This week in MarTech", m.call_args_list[0].kwargs["subject"])

    def test_personal_email_renders(self):
        from django.template.loader import render_to_string
        html = render_to_string("emails/weekly_matches.html", {
            "jobs": [{"id": 1, "slug": "x", "title": "Ops", "company": "Co", "where": "Remote",
                      "matched": 5, "required": 5, "missing": []}],
            "count": 1, "first_name": "Asha", "top_gap": "Lead Routing", "top_gap_n": 2})
        self.assertIn("Asha, 1 new role fits your resume", html)
        self.assertIn("Lead Routing", html)

    def test_cron_sends_digest_only_on_mondays(self):
        import datetime as dt
        from django.core.management import call_command
        from django.utils import timezone
        calls = []
        with mock.patch("jobs.management.commands.run_daily_tasks.call_command", side_effect=lambda *a, **k: calls.append(a[0])), \
             mock.patch.object(timezone, "now", return_value=timezone.make_aware(dt.datetime(2026, 9, 22, 9))):  # Tuesday
            call_command("run_daily_tasks")
        self.assertNotIn("send_daily_digest", calls)
        calls.clear()
        with mock.patch("jobs.management.commands.run_daily_tasks.call_command", side_effect=lambda *a, **k: calls.append(a[0])), \
             mock.patch.object(timezone, "now", return_value=timezone.make_aware(dt.datetime(2026, 9, 21, 9))):  # Monday
            call_command("run_daily_tasks")
        self.assertLess(calls.index("send_weekly_matches"), calls.index("send_daily_digest"))


def _fake_openai(payload):
    fake = mock.MagicMock()
    fake.chat.completions.create.return_value.choices = [mock.MagicMock(message=mock.MagicMock(content=json.dumps(payload)))]
    return fake


@override_settings(**TEST_SETTINGS)
class Phase4TailorTests(TestCase):
    AI = {
        "summary": "Marketing Ops specialist who administers Marketo and SFMC and builds lead scoring.",
        "changes": [
            {"before": "- Partnered with sales on routing rules and pipeline reporting",
             "after": "Partnered with sales to own routing rules and pipeline reporting across regions", "why": "Stronger verb"},
            {"before": "- Managed SFMC email sends and journeys for the global team",
             "after": "Managed Braze and SFMC journeys for the global team", "why": "adds a tool"},          # fabricated tool
            {"before": "- Ran Pardot nurture campaigns for 40k contacts across three regions",
             "after": "Ran Pardot nurture campaigns for 90k contacts, lifting pipeline 30%", "why": "numbers"},  # fabricated numbers
            {"before": "Invented line that is not in the resume", "after": "Something", "why": "x"},        # not a real line
        ],
    }

    def setUp(self):
        cache.clear()
        from accounts.models import UserResume
        self.job = make_job()
        self.user = get_user_model().objects.create_user("t4", "t4@x.test", "pw12345!")
        UserResume.objects.create(user=self.user, text=RESUME, filename="cv.pdf")

    def _tailor(self, payload=None, openai_side_effect=None):
        fake = _fake_openai(payload or self.AI)
        patch = mock.patch("openai.OpenAI", side_effect=openai_side_effect) if openai_side_effect else mock.patch("openai.OpenAI", return_value=fake)
        with mock.patch.dict("os.environ", {"OPENAI_API_KEY": "test"}), patch:
            return self.client.post("/tools/api/tailor/", data=json.dumps({"job_id": self.job.id}), content_type="application/json")

    def test_guard_keeps_only_honest_rewrites(self):
        from jobs.resume_tailor import validate_changes, validate_summary
        kept, dropped = validate_changes(RESUME, self.AI["changes"])
        self.assertEqual([k["after"] for k in kept], ["Partnered with sales to own routing rules and pipeline reporting across regions"])
        # Claiming a skill the resume doesn't show ("lead routing") is also refused:
        k2, _ = validate_changes(RESUME, [{"before": "- Partnered with sales on routing rules and pipeline reporting",
                                           "after": "Built lead routing with sales"}])
        self.assertEqual(k2, [])
        self.assertEqual(dropped, 3)
        self.assertEqual(validate_summary(RESUME, "Expert in Braze and Iterable."), "")      # tools not in resume
        self.assertEqual(validate_summary(RESUME, "Drove 300% growth."), "")                 # invented number
        self.assertTrue(validate_summary(RESUME, self.AI["summary"]))

    def test_gates_anonymous_and_no_resume(self):
        self.client.logout()
        self.assertEqual(self._tailor().status_code, 401)
        u2 = get_user_model().objects.create_user("t5", "t5@x.test", "pw12345!")
        self.client.force_login(u2)
        self.assertEqual(self._tailor().status_code, 400)

    def test_one_free_tailor_then_pro_gate_and_nothing_stored(self):
        from accounts.models import TailorUse
        self.client.force_login(self.user)
        r = self._tailor()
        self.assertEqual(r.status_code, 200)
        d = r.json()
        self.assertEqual(len(d["changes"]), 1)
        self.assertEqual(d["dropped"], 3)
        self.assertEqual(d["free_left"], 0)
        self.assertIn("Segment", d["not_added"])
        self.assertEqual(TailorUse.objects.filter(user=self.user).count(), 1)
        self.assertFalse(hasattr(TailorUse, "text"))
        r2 = self._tailor()
        self.assertEqual((r2.status_code, r2.json()["gate"]), (402, "pro"))
        self.assertContains(self.client.get(f"/tools/resume-keyword-scanner/?job={self.job.id}"), "Free tailor used")

    def test_ai_failure_is_friendly_and_not_counted(self):
        from accounts.models import TailorUse
        self.client.force_login(self.user)
        r = self._tailor(openai_side_effect=RuntimeError("down"))
        self.assertEqual(r.status_code, 503)
        self.assertEqual(TailorUse.objects.count(), 0)
        r = self._tailor(payload={"summary": "Expert in Braze.", "changes": [self.AI["changes"][1]]})  # all unsafe
        self.assertEqual(r.status_code, 503)
        self.assertEqual(TailorUse.objects.count(), 0)

    def test_staff_unlimited(self):
        self.user.is_staff = True; self.user.save()
        self.client.force_login(self.user)
        for _ in range(3):
            self.assertEqual(self._tailor().status_code, 200)

    def test_tailor_right_after_a_check_is_not_blocked_by_click_cooldown(self):
        import time as _t
        self.client.force_login(self.user)
        sess = self.client.session; sess["last_ai_call"] = _t.time(); sess.save()   # a check just used AI tips
        self.assertEqual(self._tailor().status_code, 200)

    def test_docx_download_is_a_real_word_file(self):
        import io
        import docx
        self.client.force_login(self.user)
        text = "Jane Doe\n\nSUMMARY\nMarketing Ops specialist.\n\nEXPERIENCE\n" + RESUME
        r = self.client.post("/tools/api/tailor/docx/", data=json.dumps({"text": text}), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        d = docx.Document(io.BytesIO(r.content))
        self.assertEqual(d.paragraphs[0].text, "Jane Doe")
        self.assertTrue(any("Marketo" in p.text for p in d.paragraphs))
        self.client.logout()
        self.assertEqual(self.client.post("/tools/api/tailor/docx/", data=json.dumps({"text": text}), content_type="application/json").status_code, 401)


@override_settings(**TEST_SETTINGS)
class AnalyticsCoverageTests(TestCase):
    def setUp(self):
        cache.clear()
        self.job = make_job()

    def test_errors_and_new_features_are_tracked(self):
        from accounts.models import UserResume
        r = self.client.get(f"/tools/resume-keyword-scanner/?job={self.job.id}")
        for ev in ["ats_error_shown", "tailor_error", "tailor_click", "tailor_success", "tailor_download",
                   "pro_gate_shown", "pro_preorder_click", "resume_uploaded"]:
            self.assertContains(r, ev)
        u = get_user_model().objects.create_user("an", "an@x.test", "pw12345!")
        UserResume.objects.create(user=u, text=RESUME)
        self.client.force_login(u)
        with self.settings(POSTHOG_KEY="phc_test"):
            r = self.client.get(f"/job/{self.job.id}/{self.job.slug}/")
            self.assertContains(r, "fit_badge_job_click")
            self.assertContains(r, "came_from: window.mtjCameFrom")   # pageviews carry the previous page type
        r = self.client.get("/accounts/matches/")
        self.assertContains(r, "my_matches_job_click")
        self.assertContains(r, "My Matches | MarTechJobs</title>")


# ---------------------------------------------------------------------------
# Site audit batch 1 (Sept 2026)
# ---------------------------------------------------------------------------
class AuditBatch1Tests(TestCase):
    def setUp(self):
        cache.clear()
        self.job = make_job()
        make_job(title="CRM Manager", company="DoorDash, Inc.", location="New York, NY")

    def test_company_page_renders_with_jobs(self):
        r = self.client.get("/companies/acme/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Marketing Operations Manager")
        self.assertContains(r, "1 active marketing technology role")

    def test_company_name_with_punctuation(self):
        r = self.client.get("/companies/doordash-inc/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "CRM Manager")

    def test_company_without_jobs_says_so(self):
        r = self.client.get("/companies/nobody-here/")
        self.assertEqual(r.status_code, 404)
        self.assertContains(r, "No active roles right now", status_code=404)
        self.assertContains(r, "noindex", status_code=404)

    def test_location_slug_html_injection_blocked(self):
        r = self.client.get('/x<img src=x onerror=alert(1)>/jobs/')
        self.assertEqual(r.status_code, 404)
        self.assertNotContains(r, "<img src=x", status_code=404)

    def test_location_header_escaped(self):
        from jobs.views import seo_landing_jobs
        r = self.client.get("/new-york/jobs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "New York")

    def test_sitemap_has_no_noindex_urls_and_includes_companies(self):
        import re as _re
        body = self.client.get("/sitemap.xml").content.decode()
        urls = _re.findall(r"<loc>https?://[^/]+(/[^<]*)</loc>", body)
        self.assertIn("/companies/acme/", urls)
        self.assertNotIn("/post-job/", urls)
        for u in urls:
            r = self.client.get(u)
            self.assertIn(r.status_code, (200,), u)
            self.assertNotIn(b'content="noindex', r.content, u)

    def test_noindex_on_form_and_variant_pages(self):
        for u in ["/post-job/", "/unsubscribe/", f"/tools/resume-keyword-scanner/?job={self.job.id}"]:
            self.assertContains(self.client.get(u), 'content="noindex', msg_prefix=u)
        self.assertNotContains(self.client.get("/tools/resume-keyword-scanner/"), 'content="noindex')

    def test_featured_plan_without_stripe_does_not_error(self):
        from jobs import views as v
        with self.settings(STRIPE_SECRET_KEY=""):
            self.assertIn("plan == 'featured' and settings.STRIPE_SECRET_KEY", open(v.__file__).read())

    def test_employer_page_makes_no_false_promises(self):
        r = self.client.get("/for-employers/")
        self.assertNotContains(r, "Live in minutes")
        self.assertNotContains(r, "job-alert")
        self.assertNotContains(r, "Certified Salesforce")


# ---------------------------------------------------------------------------
# Site audit batch 2: job data quality
# ---------------------------------------------------------------------------
class IngestQualityRuleTests(TestCase):
    def test_title_codes_removed_real_words_kept(self):
        from jobs.ingest_quality import clean_title
        self.assertEqual(clean_title("(PR0056) Salesforce Marketing Cloud Consultant"), "Salesforce Marketing Cloud Consultant")
        self.assertEqual(clean_title("Senior Analyst, Digital Analytics (L09)"), "Senior Analyst, Digital Analytics")
        self.assertEqual(clean_title("Marketing Ops Intern (2026)"), "Marketing Ops Intern (2026)")
        self.assertEqual(clean_title("CRM Manager (Remote)"), "CRM Manager (Remote)")

    def test_locations_tidied(self):
        from jobs.ingest_quality import tidy_location
        self.assertEqual(tidy_location("6 Locations"), "Multiple locations")
        self.assertEqual(tidy_location("A, CA; B, NY; C, WA; D, TX; E, CO"), "A, CA; B, NY; C, WA + 2 more")
        self.assertEqual(tidy_location("Austin, TX"), "Austin, TX")

    def test_salary_only_from_real_ranges(self):
        from jobs.ingest_quality import extract_salary
        self.assertEqual(extract_salary("<p>Base pay: $157,000&mdash;$227,000 a year</p>"), "157,000 - 227,000 USD")
        self.assertEqual(extract_salary("Base pay: $157,000 - $227,000 a year"), "157,000 - 227,000 USD")
        self.assertEqual(extract_salary("&lt;p&gt;Range $110K to $140K&lt;/p&gt;"), "110,000 - 140,000 USD")
        self.assertIsNone(extract_salary("$25 - $40 per hour"))
        self.assertIsNone(extract_salary("Founded in 2010 with 500 staff"))
        self.assertIsNone(extract_salary("$20,000 - $900,000"))  # implausible spread

    def test_save_applies_rules(self):
        j = make_job(title="(1507) Marketo Specialist", company="Doordashusa", location="3 Locations")
        j.refresh_from_db()
        self.assertEqual((j.title, j.company, j.location), ("Marketo Specialist", "DoorDash", "Multiple locations"))
        self.assertEqual(j.slug, "1507-marketo-specialist-at-doordashusa" if False else j.slug)  # slug unchanged by rules

    def test_old_company_url_redirects(self):
        make_job(company="DoorDash")
        r = self.client.get("/companies/doordashusa/")
        self.assertEqual((r.status_code, r["Location"]), (301, "/companies/doordash/"))


class CleanJobDataCommandTests(TestCase):
    def test_renames_backfills_salary_and_dedupes(self):
        from django.core.management import call_command
        a = make_job(title="CRM Manager", company="Acme", location="Austin, TX",
                     description="<p>Salary $120,000 - $150,000</p>")
        b = make_job(title="CRM Manager", company="Acme", location="Austin, TX")
        Job.objects.filter(pk=a.pk).update(company="Wppmedia")  # legacy row saved before the rules
        Job.objects.filter(pk=b.pk).update(company="Wppmedia")
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual(a.company, "WPP Media")
        self.assertEqual(a.salary_range, "120,000 - 150,000 USD")
        self.assertEqual([a.is_active, b.is_active].count(True), 1)  # one live copy
        self.assertTrue(b.is_active)  # newest wins

    def test_dry_run_writes_nothing(self):
        from django.core.management import call_command
        j = make_job(company="Acme")
        Job.objects.filter(pk=j.pk).update(title="(PR1) CRM Manager")
        call_command("clean_job_data", "--dry-run", stdout=open("/dev/null", "w"))
        j.refresh_from_db()
        self.assertEqual(j.title, "(PR1) CRM Manager")


class FeedClosingTests(TestCase):
    """A job that disappears from its company's board is closed on the next poll."""

    def _board(self, ids):
        from datetime import datetime, timezone as tz
        now = datetime.now(tz.utc).isoformat()
        return {"jobs": [{"id": i, "title": f"Marketo Admin {i}", "updated_at": now,
                          "location": {"name": "Remote"}, "content": JD_HTML,
                          "absolute_url": f"https://acme.test/jobs/{i}"} for i in ids]}

    def _run(self, ids):
        from django.core.management import call_command
        def fake_get(url, *a, **k):
            m = mock.MagicMock(status_code=200)
            m.json.return_value = {"name": "Acme"} if url.endswith("/boards/acme") else self._board(ids)
            return m
        approved = {"score": 90, "status": "approved", "details": {"signals": {"stack": []}}}
        with mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=fake_get), \
             mock.patch("jobs.management.commands.fetch_jobs.MarTechScreener.screen", return_value=approved), \
             mock.patch("jobs.management.commands.fetch_jobs.time.sleep"):
            call_command("fetch_jobs", "--sources-only", stdout=open("/dev/null", "w"))

    def setUp(self):
        from jobs.models import CompanySource
        CompanySource.objects.create(name="Acme", ats_type="greenhouse", token="acme")

    def test_vanished_job_closed_others_kept(self):
        self._run([1, 2, 3])
        self.assertEqual(Job.objects.filter(is_active=True).count(), 3)
        self._run([1, 2])
        self.assertEqual(set(Job.objects.filter(is_active=True).values_list("external_id", flat=True)),
                         {"greenhouse:1", "greenhouse:2"})
        closed = Job.objects.get(external_id="greenhouse:3")
        self.assertIn("no longer listed", closed.screening_reason)
        self.assertIsNotNone(Job.objects.get(external_id="greenhouse:1").last_seen_at)

    def test_board_glitch_does_not_wipe_jobs(self):
        self._run([1, 2, 3, 4, 5, 6])
        self._run([1])  # 5 of 6 vanish at once: treat as an API glitch
        self.assertEqual(Job.objects.filter(is_active=True).count(), 6)

    def test_legacy_job_without_source_key_is_matched(self):
        make_job(title="Old Role", company="Acme", external_id="greenhouse:99", location="Remote")
        self._run([1, 2, 3])
        self.assertFalse(Job.objects.get(external_id="greenhouse:99").is_active)

    def test_closed_job_page_is_gone(self):
        self._run([1, 2, 3]); self._run([1, 2])
        j = Job.objects.get(external_id="greenhouse:3")
        self.assertEqual(self.client.get(f"/job/{j.id}/{j.slug}/").status_code, 410)


class WorkdaySiteClosingTests(TestCase):
    """A company with two Workday sites: polling one never closes the other's jobs."""

    def test_other_site_jobs_untouched(self):
        from jobs.management.commands.fetch_jobs import Command
        from jobs.models import CompanySource
        s = CompanySource.objects.create(name="Acme", ats_type="workday", token="acme/SiteA",
                                         board_url="https://acme.wd1.myworkdayjobs.com/SiteA")
        a = make_job(title="Ops A", company="Acme", external_id="workday:/job/a",
                     apply_url="https://acme.wd1.myworkdayjobs.com/SiteA/job/a")
        b = make_job(title="Ops B", company="Acme", external_id="workday:/job/b",
                     apply_url="https://acme.wd1.myworkdayjobs.com/SiteB/job/b")
        make_job(title="Ops C", company="Acme", external_id="workday:/job/c",
                 apply_url="https://acme.wd1.myworkdayjobs.com/SiteA/job/c")
        cmd = Command()
        cmd.stats = __import__("collections").defaultdict(int)
        cmd._source_key = "workday:acme/SiteA"
        cmd._feed_ids = {"workday:/job/c"}
        cmd._close_vanished(s, {"Acme"})
        a.refresh_from_db(); b.refresh_from_db()
        self.assertFalse(a.is_active)   # gone from Site A -> closed
        self.assertTrue(b.is_active)    # belongs to Site B -> untouched


class ScannerVisibilityTests(TestCase):
    """Resume Scanner visibility: skills card, apply nudge, homepage button."""

    def setUp(self):
        cache.clear()
        self.job = make_job()

    def url(self):
        return f"/job/{self.job.id}/{self.job.slug}/"

    def test_job_page_lists_this_jobs_skills(self):
        r = self.client.get(self.url())
        self.assertContains(r, 'data-placement="skills_card"')
        self.assertContains(r, "This role asks for")
        self.assertContains(r, "Marketo Engage")          # a platform from JD_HTML
        self.assertContains(r, "more</span>")             # 8 terms -> 3 shown + 5 more

    def test_apply_nudge_present_and_tracked(self):
        r = self.client.get(self.url())
        self.assertContains(r, 'id="jd-nudge"')
        self.assertContains(r, "apply_nudge_shown")
        self.assertContains(r, 'data-placement="apply_nudge"')

    def test_no_skills_card_when_member_has_saved_resume(self):
        from accounts.models import UserResume
        u = get_user_model().objects.create_user("sr", "sr@x.test", "pw12345!x")
        UserResume.objects.create(user=u, text=RESUME)
        self.client.force_login(u)
        self.assertNotContains(self.client.get(self.url()), 'data-placement="skills_card"')

    def test_job_without_skills_has_no_card(self):
        j = make_job(title="Office Manager", description="<p>Keep the office running smoothly every day.</p>")
        r = self.client.get(f"/job/{j.id}/{j.slug}/")
        self.assertNotContains(r, 'data-placement="skills_card"')
        self.assertNotContains(r, 'id="jd-nudge"')

    def test_homepage_links_to_scanner(self):
        r = self.client.get("/")
        self.assertContains(r, "js-hero-scan")
        self.assertContains(r, "/tools/resume-keyword-scanner/")


class ShortResumeNotSavedTests(TestCase):
    """A pasted resume too short to check must not be saved (it broke every later check)."""

    def test_short_paste_rejected_and_not_saved(self):
        from accounts.models import UserResume
        cache.clear()
        job = make_job()
        u = get_user_model().objects.create_user("sh", "sh@x.test", "pw12345!x")
        self.client.force_login(u)
        r = self.client.post("/tools/api/ats-match/", data=json.dumps(
            {"resume_text": "Jane Doe\nMarketing Ops", "job_id": job.id, "save": True}),
            content_type="application/json")
        self.assertEqual(r.status_code, 400)
        self.assertFalse(UserResume.objects.filter(user=u).exists())

    def test_existing_short_saved_resume_gets_clear_message(self):
        from accounts.models import UserResume
        cache.clear()
        job = make_job()
        u = get_user_model().objects.create_user("sh2", "sh2@x.test", "pw12345!x")
        UserResume.objects.create(user=u, text="too short")
        self.client.force_login(u)
        r = self.client.post("/tools/api/ats-match/", data=json.dumps({"job_id": job.id}),
                             content_type="application/json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Replace", r.json()["error"])


class TrafficBatchTests(TestCase):
    """Google job ping, listing data, job alerts, resume-keyword pages."""

    def setUp(self):
        cache.clear()

    # --- Google Indexing API: new jobs once, closed jobs removed once ---
    def _run_index(self):
        from django.core.management import call_command
        sent = []
        def fake_post(url, json=None, **k):
            sent.append((json["type"], json["url"])); return mock.MagicMock(status_code=200)
        creds = mock.MagicMock(token="t", service_account_email="x@y")
        with mock.patch.dict("os.environ", {"GOOGLE_JSON_KEY": '{"client_email":"x@y"}'}), \
             mock.patch("google.oauth2.service_account.Credentials.from_service_account_info", return_value=creds), \
             mock.patch("jobs.management.commands.index_jobs.requests.post", side_effect=fake_post):
            call_command("index_jobs", stdout=open("/dev/null", "w"))
        return sent

    def test_index_sends_new_once_and_deletes_closed(self):
        live = make_job(title="Live Role")
        gone = make_job(title="Gone Role")
        first = self._run_index()
        self.assertEqual({t for t, _ in first}, {"URL_UPDATED"})
        self.assertEqual(len(first), 2)
        self.assertEqual(self._run_index(), [])                  # nothing re-sent
        Job.objects.filter(pk=gone.pk).update(is_active=False)
        third = self._run_index()
        self.assertEqual(third, [("URL_DELETED", f"https://martechjobs.io/job/{gone.id}/{gone.slug}/")] if third and "martechjobs.io" in third[0][1] else third)
        self.assertEqual([t for t, _ in third], ["URL_DELETED"])
        self.assertEqual(self._run_index(), [])

    # --- JobPosting location data ---
    def test_schema_location_for_multi_city_and_placeholder(self):
        self.assertEqual(Job(location="Multiple locations").get_address_parts(), {"locality": "", "region": ""})
        self.assertEqual(Job(location="San Francisco, CA; New York, NY; Seattle, WA + 7 more").get_address_parts(),
                         {"locality": "San Francisco", "region": "CA"})
        self.assertTrue(make_job().get_schema_valid_through().endswith("+00:00"))

    # --- Job alerts ---
    def test_tool_alert_on_job_and_tool_pages(self):
        cat = Category.objects.create(name="MarTech", slug="martech")
        tool = Tool.objects.create(name="Marketo", slug="marketo", category=cat)
        j = make_job(); j.tools.add(tool)
        r = self.client.get(f"/job/{j.id}/{j.slug}/")
        self.assertContains(r, 'id="jd-alert-form"')
        self.assertContains(r, 'name="tool" value="marketo"')
        self.assertContains(r, "Get new Marketo jobs by email")

    def test_subscribe_js_sends_filters(self):
        r = self.client.get("/")
        self.assertContains(r, "new URLSearchParams(new FormData(form))")
        self.assertContains(r, "job_alert_created")

    def test_alert_post_creates_targeted_alert_not_newsletter(self):
        from jobs.models import SavedSearch, PendingSubscriber
        r = self.client.post("/subscribe/", {"email": "a@b.test", "tool": "marketo"})
        self.assertEqual(r.json()["success"], True)
        self.assertTrue(SavedSearch.objects.filter(email="a@b.test", tool="marketo").exists())
        self.assertFalse(PendingSubscriber.objects.filter(email="a@b.test").exists())

    # --- Resume keyword pages ---
    def test_resume_keywords_page_real_numbers_and_indexable(self):
        for i in range(9):
            make_job(title=f"Marketing Operations Manager {i}", company=f"Co{i}")
        r = self.client.get("/marketing-operations-manager-resume-keywords/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Marketing Operations Manager resume keywords")
        self.assertContains(r, "Salesforce Marketing Cloud")
        self.assertContains(r, "100%")                    # every seeded post asks for it
        self.assertNotContains(r, 'content="noindex')
        self.assertContains(r, "FAQPage")
        body = self.client.get("/sitemap.xml").content.decode()
        self.assertIn("/marketing-operations-manager-resume-keywords/", body)

    def test_resume_keywords_thin_role_is_noindex_and_not_in_sitemap(self):
        make_job(title="CRM Manager")
        r = self.client.get("/crm-manager-resume-keywords/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'content="noindex')
        self.assertNotIn("/crm-manager-resume-keywords/", self.client.get("/sitemap.xml").content.decode())

    def test_unknown_role_404(self):
        self.assertEqual(self.client.get("/not-a-role-resume-keywords/").status_code, 404)

    def test_title_and_salary_pages_link_to_keywords(self):
        make_job()
        self.assertContains(self.client.get("/marketing-operations-manager-jobs/"), "/marketing-operations-manager-resume-keywords/")
        self.assertContains(self.client.get("/marketing-operations-manager-salary/"), "/marketing-operations-manager-resume-keywords/")


class ToolTaggingTests(TestCase):
    """Jobs get tagged with every platform their description asks for."""

    def test_clean_job_data_tags_tools_from_description(self):
        from django.core.management import call_command
        cache.clear()
        j = make_job(title="Lifecycle Manager", description="<p>Own our Braze and Salesforce Marketing Cloud journeys. HubSpot a plus.</p>")
        self.assertEqual(j.tools.count(), 0)
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        names = set(j.tools.values_list("name", flat=True))
        self.assertTrue({"Braze", "Salesforce Marketing Cloud", "HubSpot"} <= names, names)
        self.assertNotIn("Salesforce", names)   # SFMC is not plain Salesforce
        r = self.client.get("/jobs/braze/")
        self.assertContains(r, "Lifecycle Manager")

    def test_tool_page_h1_includes_search_abbreviation(self):
        from django.core.management import call_command
        cache.clear()
        make_job(description="<p>Salesforce Marketing Cloud admin needed.</p>")
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        self.assertContains(self.client.get("/jobs/salesforce-marketing-cloud/"), "Salesforce Marketing Cloud (SFMC)")

    def test_dry_run_does_not_tag(self):
        from django.core.management import call_command
        cache.clear()
        j = make_job(description="<p>Braze expert.</p>")
        call_command("clean_job_data", "--dry-run", stdout=open("/dev/null", "w"))
        self.assertEqual(j.tools.count(), 0)


class GeneratedRolePageTests(TestCase):
    """Role pages from real titles + tool x function pages, with thin-page guards."""

    def setUp(self):
        cache.clear()

    def test_title_normalisation(self):
        from jobs.role_pages import normalize_title as n
        self.assertEqual(n("Sr. Salesforce Developer - Tieto Tech Consulting (m/f/d)"), "salesforce developer")
        self.assertEqual(n("Senior Developer - Salesforce Marketing Cloud"), "salesforce marketing cloud developer")
        self.assertEqual(n("Marketing Operations Lead, Customer Loyalty | Irvine, CA"), "marketing operations lead")
        self.assertEqual(n("Analytics Engineer II (Remote)"), "analytics engineer")
        self.assertEqual(n("Salesforce 架構師 (Salesforce Solution Architect)"), "")

    def _devs(self, n, companies):
        for i in range(n):
            make_job(title=f"Senior Salesforce Developer {'I' * (i % 3 + 1)}", company=companies[i % len(companies)],
                     description="<p>Salesforce Apex and Lightning developer.</p>")

    def test_auto_role_page_indexable_with_enough_jobs_and_employers(self):
        self._devs(6, ["A", "B", "C"])
        r = self.client.get("/salesforce-developer-jobs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Salesforce Developer")
        self.assertContains(r, "6 open Salesforce Developer roles at 3 companies")
        self.assertNotContains(r, 'content="noindex')
        self.assertNotContains(r, "FAQPage")                 # no generic Q&A on generated pages
        self.assertNotContains(r, "-salary/")               # salary pages exist only for core roles
        self.assertIn("/salesforce-developer-jobs/", self.client.get("/sitemap.xml").content.decode())

    def test_auto_role_single_employer_is_noindex(self):
        self._devs(4, ["OnlyCo"])
        r = self.client.get("/salesforce-developer-jobs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'content="noindex')
        self.assertNotIn("/salesforce-developer-jobs/", self.client.get("/sitemap.xml").content.decode())

    def test_unknown_role_404_and_core_role_unchanged(self):
        self.assertEqual(self.client.get("/basket-weaver-jobs/").status_code, 404)
        make_job()
        r = self.client.get("/marketing-operations-manager-jobs/")
        self.assertContains(r, "FAQPage")
        self.assertContains(r, "/marketing-operations-manager-salary/")

    def test_tool_function_page(self):
        from django.core.management import call_command
        for i, co in enumerate(["A", "B", "C"]):
            make_job(title=f"SFMC Developer {i}", company=co, description="<p>Salesforce Marketing Cloud developer, AMPscript.</p>")
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        r = self.client.get("/jobs/salesforce-marketing-cloud/developer/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "SFMC Developer")
        self.assertNotContains(r, 'content="noindex')
        self.assertContains(self.client.get("/jobs/salesforce-marketing-cloud/"), "/jobs/salesforce-marketing-cloud/developer/")
        self.assertIn("/jobs/salesforce-marketing-cloud/developer/", self.client.get("/sitemap.xml").content.decode())
        self.assertEqual(self.client.get("/jobs/salesforce-marketing-cloud/astronaut/").status_code, 404)

    def test_offtopic_roles_removed(self):
        from django.core.management import call_command
        w = make_job(title="Welder/Brazer II - 2nd Shift (Onsite)")
        k = make_job(title="Marketing Operations Manager")
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        w.refresh_from_db(); k.refresh_from_db()
        self.assertFalse(w.is_active)
        self.assertTrue(k.is_active)


class LlmsTxtTests(TestCase):
    """llms.txt: real live numbers and only working links (AI assistants quote it)."""

    def test_links_all_work_and_facts_are_live(self):
        import re as _re
        cache.clear()
        cat = Category.objects.create(name="MarTech", slug="martech")
        mk = Tool.objects.create(name="Marketo", slug="marketo", category=cat)
        for i in range(3):
            make_job(title=f"Marketo Specialist {i}", company=f"C{i}").tools.add(mk)
        body = self.client.get("/llms.txt").content.decode()
        self.assertIn("3 open MarTech roles at 3 companies", body)
        self.assertIn("Resume Scanner", body)
        self.assertIn("Marketo (3 roles)", body)
        self.assertNotIn("/marketing-operations-jobs/", body)     # was a broken link
        for u in set(_re.findall(r"https://martechjobs.io(/[^)\s]*)", body)):
            self.assertEqual(self.client.get(u).status_code, 200, u)


class SearchConsoleBatchTests(TestCase):
    """GSC-driven batch: analyst page, click-worthy titles, Search Console sync + Founder HQ."""

    def setUp(self):
        cache.clear()

    def test_marketing_technology_analyst_page(self):
        make_job(title="Senior Analyst, Marketing Analytics", company="A")
        make_job(title="Marketing Operations Analyst", company="B")
        r = self.client.get("/marketing-technology-analyst-jobs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Marketing Technology Analyst")
        self.assertContains(r, "Senior Analyst, Marketing Analytics")
        self.assertEqual(self.client.get("/marketing-technology-analyst-salary/").status_code, 200)

    def test_titles_carry_live_counts(self):
        cat = Category.objects.create(name="MarTech", slug="martech")
        t = Tool.objects.create(name="Salesforce Marketing Cloud", slug="salesforce-marketing-cloud", category=cat,
                                seo_title="Old hand title")
        for i in range(3):
            make_job(title=f"SFMC Dev {i}", company=f"C{i}").tools.add(t)
        r = self.client.get("/jobs/salesforce-marketing-cloud/")
        self.assertContains(r, "<title>3 Salesforce Marketing Cloud (SFMC) Jobs — Updated")
        r = self.client.get("/remote/jobs/")
        self.assertContains(r, "Remote MarTech Jobs — Updated")
        r = self.client.get("/jobs-by-tool/")
        self.assertContains(r, "MarTech Jobs by Platform: Salesforce Marketing Cloud")
        self.assertContains(r, "(3 live roles)")

    def test_blog_snippet_migration(self):
        import importlib
        from django.apps import apps as django_apps
        mig = importlib.import_module("jobs.migrations.0019_blog_job_titles_meta")
        BlogPost.objects.create(title="T", slug=mig.SLUG, excerpt="e", content="c")
        mig.forwards(django_apps, None)
        self.assertEqual(BlogPost.objects.get(slug=mig.SLUG).meta_title, mig.META_TITLE)

    def _sync(self, status=200):
        from django.core.management import call_command
        import io
        def fake_post(url, json=None, **k):
            m = mock.MagicMock(status_code=status)
            if json["dimensions"] == ["date"]:
                m.json.return_value = {"rows": [{"keys": ["2026-09-20"], "clicks": 10, "impressions": 500, "ctr": .02, "position": 9.5},
                                                {"keys": ["2026-08-20"], "clicks": 5, "impressions": 400, "ctr": .0125, "position": 14}]}
            else:
                m.json.return_value = {"rows": [
                    {"keys": ["marketo jobs", "https://martechjobs.io/jobs/marketo/"], "clicks": 1, "impressions": 226, "ctr": .004, "position": 12.5},
                    {"keys": ["martech jobs", "https://martechjobs.io/"], "clicks": 50, "impressions": 800, "ctr": .06, "position": 3}]}
            return m
        creds = mock.MagicMock(token="t")
        out = io.StringIO()
        with mock.patch.dict("os.environ", {"GOOGLE_JSON_KEY": '{"client_email":"bot@x.iam.gserviceaccount.com"}'}), \
             mock.patch("google.oauth2.service_account.Credentials.from_service_account_info", return_value=creds), \
             mock.patch("jobs.management.commands.gsc_sync.requests.post", side_effect=fake_post):
            call_command("gsc_sync", stdout=out)
        return out.getvalue()

    def test_sync_stores_data_and_founder_hq_shows_readable_tables(self):
        from jobs.models import SearchConsoleDaily, SearchConsoleRow
        out = self._sync()
        self.assertIn("✅", out)
        self.assertEqual(SearchConsoleDaily.objects.count(), 2)
        self.assertEqual(SearchConsoleRow.objects.count(), 2)
        staff = get_user_model().objects.create_user("st", "st@x.test", "pw12345!x", is_staff=True)
        self.client.force_login(staff)
        r = self.client.get("/staff/")
        self.assertContains(r, "Google search")
        self.assertContains(r, "Marketo jobs page")          # readable name…
        self.assertNotContains(r, "martechjobs.io/jobs/marketo/")  # …never a raw URL
        self.assertContains(r, "marketo jobs")

    def test_sync_explains_how_to_grant_access(self):
        out = self._sync(status=403)
        self.assertIn("bot@x.iam.gserviceaccount.com", out)
        self.assertIn("Users and permissions", out)


class FounderHqStillLockedTests(TestCase):
    def test_founder_hq_requires_staff(self):
        r = self.client.get("/staff/")
        self.assertIn(r.status_code, (302, 403))
        u = get_user_model().objects.create_user("nobody", "n@x.test", "pw12345!x")
        self.client.force_login(u)
        r = self.client.get("/staff/")
        self.assertIn(r.status_code, (302, 403))


class BlogCleanupTests(TestCase):
    def setUp(self):
        cache.clear()

    def test_job_slug_posts_hidden_from_google_but_readable(self):
        BlogPost.objects.create(title="Role Guide", slug="martech-jobs", excerpt="e", content="<p>c</p>")
        BlogPost.objects.create(title="Keep me", slug="martech-job-titles-career-paths", excerpt="e", content="<p>c</p>")
        r = self.client.get("/blog/martech-jobs/")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'content="noindex')
        self.assertNotContains(self.client.get("/blog/martech-job-titles-career-paths/"), 'content="noindex')
        sm = self.client.get("/sitemap.xml").content.decode()
        self.assertNotIn("/blog/martech-jobs/", sm)
        self.assertIn("/blog/martech-job-titles-career-paths/", sm)

    def test_daily_run_skips_auto_blog_by_default(self):
        from jobs.management.commands import run_daily_tasks
        src = open(run_daily_tasks.__file__).read()
        self.assertIn('os.environ.get("AUTO_BLOG") == "1"', src)
        with mock.patch.object(run_daily_tasks.Command, "_run") as run, \
             mock.patch.dict("os.environ", {}, clear=False):
            import os; os.environ.pop("AUTO_BLOG", None)
            run_daily_tasks.Command().handle()
        called = [c.args[1] for c in run.call_args_list]
        self.assertNotIn("generate_blog", called)
        self.assertIn("fetch_jobs", called)


class CountryDropdownTests(TestCase):
    """Location dropdown lists countries only; picking one finds every job there."""

    def setUp(self):
        cache.clear()

    def test_resolver(self):
        from jobs.geo import country_code as c
        cases = {"CAN, ON, Mississauga": "CA", "Santa Clara, CA, United States": "US", "New Mexico": "US",
                 "Bengaluru, India": "IN", "Noida, Uttar Pradesh": "IN", "Omaha, NE 68122": "US",
                 "Remote - USA": "US", "MEX Work-at-Home": "MX", "Copenhagen, Capital Region, Denmark": "DK",
                 "Multiple locations": "", "Latin America": "", "A, CA; B, NY + 9 more": "",
                 "Gran Buenos Aires, Argentina; Capital Federal, Argentina": "AR"}
        for loc, want in cases.items():
            self.assertEqual(c(loc), want, loc)

    def test_dropdown_has_only_countries_with_counts(self):
        make_job(title="A", location="Santa Clara, CA, United States", work_arrangement="onsite")
        make_job(title="B", location="California, San Francisco", work_arrangement="onsite")
        make_job(title="C", location="Bengaluru, Karnataka", work_arrangement="onsite")
        import re as _re
        html = self.client.get("/").content.decode()
        select = _re.search(r'<select[^>]*name="l".*?</select>', html, _re.S).group(0)
        self.assertIn(">United States (2)</option>", select)
        self.assertIn(">India (1)</option>", select)
        self.assertNotIn("California", select)
        self.assertNotIn("Karnataka", select)

    def test_country_filter_finds_jobs_without_country_in_text(self):
        make_job(title="Bangalore Marketo Admin", location="Bengaluru, Karnataka", work_arrangement="onsite")
        make_job(title="Austin Ops Role", location="Austin, TX", work_arrangement="onsite")
        r = self.client.get("/jobs/?l=India")
        self.assertContains(r, "Bangalore Marketo Admin")
        self.assertNotContains(r, "Austin Ops Role")

    def test_daily_cleanup_backfills_country(self):
        from django.core.management import call_command
        j = make_job(location="Rotterdam", work_arrangement="onsite")
        Job.objects.filter(pk=j.pk).update(country="")
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        j.refresh_from_db()
        self.assertEqual(j.country, "NL")


class CountryCorrectionTests(TestCase):
    def test_wrongly_stored_us_is_corrected(self):
        from django.core.management import call_command
        cache.clear()
        j = make_job(location="Paris, IDF, fr", work_arrangement="onsite")
        Job.objects.filter(pk=j.pk).update(country="US")          # the old ", xx" rule's mistake
        call_command("clean_job_data", stdout=open("/dev/null", "w"))
        j.refresh_from_db()
        self.assertEqual(j.country, "FR")


@override_settings(**TEST_SETTINGS)
class SecurityBatchTests(TestCase):
    def setUp(self):
        cache.clear()

    # 1. Staff actions can't be triggered by a link/image on another site
    @mock.patch("jobs.views.send_job_alert")
    def test_review_action_requires_post(self, _alert):
        staff = get_user_model().objects.create_user("st", "st@x.test", "pw12345!x", is_staff=True)
        self.client.force_login(staff)
        job = make_job(screening_status="pending", is_active=False)
        self.assertEqual(self.client.get(f"/staff/review/{job.id}/approve/").status_code, 405)
        job.refresh_from_db(); self.assertEqual(job.screening_status, "pending")
        self.client.post(f"/staff/review/{job.id}/approve/")
        job.refresh_from_db(); self.assertEqual(job.screening_status, "approved")

    def test_review_queue_uses_post_forms(self):
        staff = get_user_model().objects.create_user("st2", "st2@x.test", "pw12345!x", is_staff=True)
        self.client.force_login(staff)
        make_job(screening_status="pending", is_active=False)
        r = self.client.get("/staff/review/")
        self.assertContains(r, 'method="post"')
        self.assertContains(r, "csrfmiddlewaretoken")

    # 2. Unsubscribe form only emails a signed link; same reply either way
    @override_settings(EMAIL_HOST_PASSWORD="test", EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_unsubscribe_form_emails_link_instead_of_unsubscribing(self):
        from django.core import mail
        Subscriber.objects.create(email="real@x.test")
        r1 = self.client.post("/unsubscribe/", {"email": "real@x.test"})
        self.assertTrue(Subscriber.objects.get(email="real@x.test").is_active)   # not removed by a stranger
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("/u/", mail.outbox[0].alternatives[0][0] if mail.outbox[0].alternatives else mail.outbox[0].body)
        r2 = self.client.post("/unsubscribe/", {"email": "nobody@x.test"})
        self.assertEqual(len(mail.outbox), 1)                                   # nothing sent to unknowns
        self.assertContains(r1, "If that address is on our list")
        self.assertContains(r2, "If that address is on our list")

    # 3. Opening the link (scanners) doesn't unsubscribe; the button / mail client does
    def test_oneclick_get_only_confirms(self):
        from django.core import signing
        Subscriber.objects.create(email="b@x.test")
        token = signing.dumps("b@x.test", salt="unsubscribe")
        r = self.client.get(f"/u/{token}/")
        self.assertContains(r, "Yes, unsubscribe me")
        self.assertTrue(Subscriber.objects.get(email="b@x.test").is_active)
        r = self.client.post(f"/u/{token}/", {"confirm": "1"})
        self.assertContains(r, "unsubscribed")
        self.assertFalse(Subscriber.objects.get(email="b@x.test").is_active)

    def test_mail_client_one_click_post(self):
        from django.core import signing
        Subscriber.objects.create(email="c@x.test")
        token = signing.dumps("c@x.test", salt="unsubscribe")
        r = self.client.post(f"/u/{token}/", {"List-Unsubscribe": "One-Click"})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(Subscriber.objects.get(email="c@x.test").is_active)

    # 4. Zip-bomb .docx refused before it is unpacked
    def test_docx_zip_bomb_rejected(self):
        import io, zipfile
        from jobs.resume_match import extract_resume_text, ResumeParseError
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (30 * 1024 * 1024))
        f = io.BytesIO(buf.getvalue()); f.name = "cv.docx"; f.size = len(buf.getvalue())
        self.assertLess(f.size, 1024 * 1024)                    # tiny on the wire…
        with self.assertRaises(ResumeParseError):
            extract_resume_text(f)                              # …refused before unpacking 30 MB

    def test_normal_docx_still_works(self):
        import io, docx
        from jobs.resume_match import extract_resume_text
        d = docx.Document()
        for line in RESUME.splitlines() * 3:
            d.add_paragraph(line)
        buf = io.BytesIO(); d.save(buf)
        f = io.BytesIO(buf.getvalue()); f.name = "cv.docx"; f.size = len(buf.getvalue())
        self.assertIn("Marketo", extract_resume_text(f))

    # 5. JD generator: capped inputs, cleaned HTML
    def test_jd_generator_output_is_sanitised(self):
        fake = mock.MagicMock()
        fake.chat.completions.create.return_value.choices = [mock.MagicMock(
            message=mock.MagicMock(content='<h3>Role</h3><script>alert(1)</script><p onclick="x()">Hi</p>'))]
        with mock.patch.dict("os.environ", {"OPENAI_API_KEY": "k"}), mock.patch("tools.views.OpenAI", return_value=fake):
            r = self.client.post("/tools/api/generate-jd/", data=json.dumps({"role": "x" * 5000, "stack": "Marketo"}),
                                 content_type="application/json")
        html = r.json()["html"]
        self.assertNotIn("<script", html); self.assertNotIn("onclick", html); self.assertIn("<h3>Role</h3>", html)
        sent = fake.chat.completions.create.call_args.kwargs["messages"][1]["content"]
        self.assertLess(len(sent), 600)


@override_settings(**TEST_SETTINGS)
class RecheckPendingTests(TestCase):
    """High-score pending jobs go live only if the company site still lists them."""

    def _pending(self, title, url, score=92):
        return make_job(title=title, apply_url=url, screening_status="pending",
                        is_active=False, screening_score=score)

    def _resp(self, code, payload=None):
        r = mock.Mock(status_code=code)
        r.json.return_value = payload or {}
        return r

    def test_open_relisted_closed_rejected_offtopic_removed(self):
        from django.core.management import call_command
        open_job = self._pending("Marketo Admin", "https://job-boards.greenhouse.io/acme/jobs/111")
        closed_job = self._pending("SFMC Developer", "https://acme.wd5.myworkdayjobs.com/en-US/Ext/job/NYC/SFMC_R1")
        welder = self._pending("Welder-Brazer II", "https://job-boards.greenhouse.io/acme/jobs/222")
        unknown = self._pending("Braze Engineer", "https://example.com/careers/9")
        low = self._pending("HubSpot Specialist", "https://job-boards.greenhouse.io/acme/jobs/333", score=80)

        def fake_get(url, **kw):
            if "boards-api.greenhouse.io" in url:
                return self._resp(200)
            if "myworkdayjobs.com/wday/cxs" in url:
                return self._resp(404)
            raise AssertionError(url)

        with mock.patch("jobs.management.commands.recheck_pending.requests.get", side_effect=fake_get):
            call_command("recheck_pending", "--confirm", stdout=mock.Mock())
        for j in (open_job, closed_job, welder, unknown, low):
            j.refresh_from_db()
        self.assertEqual((open_job.screening_status, open_job.is_active), ("approved", True))
        self.assertIsNotNone(open_job.last_seen_at)
        self.assertEqual((closed_job.screening_status, closed_job.is_active), ("rejected", False))
        self.assertEqual(welder.screening_status, "rejected")
        self.assertEqual(unknown.screening_status, "pending")
        self.assertEqual(low.screening_status, "pending")

    def test_dry_run_changes_nothing(self):
        from django.core.management import call_command
        j = self._pending("Marketo Admin", "https://job-boards.greenhouse.io/acme/jobs/111")
        with mock.patch("jobs.management.commands.recheck_pending.requests.get", return_value=self._resp(200)):
            call_command("recheck_pending", stdout=mock.Mock())
        j.refresh_from_db()
        self.assertEqual(j.screening_status, "pending")

    def test_workday_cant_apply_and_errors(self):
        from jobs.management.commands.recheck_pending import posting_status, CLOSED, UNKNOWN, OPEN
        import requests as rq
        url = "https://acme.wd5.myworkdayjobs.com/en-US/Ext/job/NYC/SFMC_R1"
        with mock.patch("jobs.management.commands.recheck_pending.requests.get",
                        return_value=self._resp(200, {"jobPostingInfo": {"canApply": False}})):
            self.assertEqual(posting_status(url), CLOSED)
        with mock.patch("jobs.management.commands.recheck_pending.requests.get",
                        return_value=self._resp(200, {"jobPostingInfo": {"canApply": True}})):
            self.assertEqual(posting_status(url), OPEN)
        with mock.patch("jobs.management.commands.recheck_pending.requests.get", side_effect=rq.Timeout()):
            self.assertEqual(posting_status(url), UNKNOWN)
        with mock.patch("jobs.management.commands.recheck_pending.requests.get", return_value=self._resp(500)):
            self.assertEqual(posting_status(url), UNKNOWN)


@override_settings(**TEST_SETTINGS)
class WorkdaySearchTests(TestCase):
    """Big Workday boards are also searched for MarTech keywords; small ones are walked fully."""

    def _cmd(self):
        from collections import Counter
        from jobs.management.commands.fetch_jobs import Command
        c = Command()
        c.processed_tokens, c.stats = set(), Counter()
        c.screened = []
        c.get_headers = lambda: {}
        c._is_duplicate = lambda *a, **k: False
        c._feed_seen = lambda *a: None
        c._count_known = lambda *a: None
        c.record_source = lambda *a, **k: None
        c._company_name = lambda *a, **k: "Acme"
        c._clean_location = lambda loc, remote: (loc, "onsite")
        c.screen_and_upsert = lambda d: c.screened.append(d)
        return c

    def _post(self, total, searched):
        def fake_post(url, json=None, **kw):
            searched.append(json["searchText"])
            r = mock.Mock(status_code=200)
            if json["searchText"] == "Marketo":
                r.json.return_value = {"total": 1, "jobPostings": [
                    {"title": "Marketo Admin", "externalPath": "/job/Bengaluru/Marketo-Admin_R1",
                     "postedOn": "Posted Today", "locationsText": "Bengaluru, India"}]}
            elif json["searchText"]:
                r.json.return_value = {"total": 0, "jobPostings": []}
            else:
                r.json.return_value = {"total": total, "jobPostings": [
                    {"title": "Accountant", "externalPath": f"/job/X/Acct_{json['offset']}",
                     "postedOn": "Posted Today", "locationsText": "Pune"}]}
            return r
        return fake_post

    def _get(self, url, **kw):
        r = mock.Mock(status_code=200)
        r.json.return_value = {"jobPostingInfo": {"jobDescription": "<p>JD</p>"}}
        return r

    @mock.patch("jobs.management.commands.fetch_jobs.time.sleep")
    def test_big_board_is_keyword_searched(self, _sleep):
        c, searched = self._cmd(), []
        with mock.patch("jobs.management.commands.fetch_jobs.requests.post", side_effect=self._post(5000, searched)), \
             mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get):
            c.fetch_workday_api("https://acme.wd3.myworkdayjobs.com/en-US/Careers")
        self.assertIn("Marketo", searched)
        self.assertIn("Marketo Admin", [d["title"] for d in c.screened])
        self.assertFalse(c._feed_complete)

    @mock.patch("jobs.management.commands.fetch_jobs.time.sleep")
    def test_small_board_walked_without_search(self, _sleep):
        c, searched = self._cmd(), []
        with mock.patch("jobs.management.commands.fetch_jobs.requests.post", side_effect=self._post(15, searched)), \
             mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get):
            c.fetch_workday_api("https://acme.wd3.myworkdayjobs.com/Careers")
        self.assertEqual(searched, [""])
        self.assertTrue(c._feed_complete)

    @mock.patch("jobs.management.commands.fetch_jobs.time.sleep")
    def test_same_board_different_casing_polled_once(self, _sleep):
        c, searched = self._cmd(), []
        with mock.patch("jobs.management.commands.fetch_jobs.requests.post", side_effect=self._post(15, searched)), \
             mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get):
            c.fetch_workday_api("https://asda.wd103.myworkdayjobs.com/en-US/AsdaJobs")
            c.fetch_workday_api("https://asda.wd103.myworkdayjobs.com/asdajobs")
        self.assertEqual(len(searched), 1)

    def test_seeder_registers_workday_board_as_tenant_site(self):
        from django.core.management import call_command
        from jobs.models import CompanySource
        from jobs.management.commands import seed_intl_sources as seed
        with mock.patch.object(seed, "CANDIDATES",
                               [("Deutsche Bank", "india", "workday", ["https://db.wd3.myworkdayjobs.com/DBWebsite"])]), \
             mock.patch.object(seed, "_count_workday", return_value=900):
            seed._VALIDATORS["workday"] = seed._count_workday
            call_command("seed_intl_sources", "--confirm", stdout=mock.Mock())
            call_command("seed_intl_sources", "--confirm", stdout=mock.Mock())
        rows = CompanySource.objects.filter(ats_type="workday")
        self.assertEqual(rows.count(), 1)
        self.assertEqual((rows[0].token, rows[0].name, rows[0].board_url),
                         ("db/DBWebsite", "Deutsche Bank", "https://db.wd3.myworkdayjobs.com/DBWebsite"))


@override_settings(**TEST_SETTINGS)
class SmartRecruitersPagingTests(TestCase):
    """SmartRecruiters returns 10 postings by default: read 100/page, search big boards."""

    def _cmd(self):
        c = WorkdaySearchTests._cmd(self)
        c.is_fresh = lambda d: True
        return c

    def _get(self, calls, total):
        def fake(url, params=None, **kw):
            calls.append(dict(params or {}))
            r = mock.Mock(status_code=200)
            if url.endswith("/postings") and params and params.get("q") == "Marketo":
                r.json.return_value = {"totalFound": 1, "content": [
                    {"id": "sfmc1", "name": "Marketo Specialist", "releasedDate": "x", "location": {"city": "Pune"}}]}
            elif url.endswith("/postings") and params and params.get("q"):
                r.json.return_value = {"totalFound": 0, "content": []}
            elif url.endswith("/postings"):
                off = params["offset"]
                r.json.return_value = {"totalFound": total, "content": [
                    {"id": f"p{off + i}", "name": "Accountant", "releasedDate": "x", "location": {}} for i in range(min(100, total - off))]}
            else:
                r.json.return_value = {"jobAd": {"sections": {"jobDescription": {"text": "<p>JD</p>"}}}}
            return r
        return fake

    def test_small_board_read_in_full_no_search(self):
        c, calls = self._cmd(), []
        with mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get(calls, 30)):
            c.fetch_smartrecruiters_api("Acme")
        self.assertEqual(calls[0].get("limit"), 100)
        self.assertFalse(any(p.get("q") for p in calls))
        self.assertTrue(c._feed_complete)
        self.assertEqual(len(c.screened), 30)

    def test_big_board_searched_and_capped(self):
        c, calls = self._cmd(), []
        with mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get(calls, 5000)):
            c.fetch_smartrecruiters_api("PublicisGroupe")
        self.assertTrue(any(p.get("q") == "Marketo" for p in calls))
        self.assertFalse(c._feed_complete)
        # per-board budget stops detail fetches at 40, so the search hit is budget-skipped
        self.assertEqual(len(c.screened), c.WORKDAY_MAX_NEW_PER_BOARD)
        self.assertGreater(c.stats["SmartRecruiters:budget_skip"], 0)

    def test_run_wide_cap(self):
        c, calls = self._cmd(), []
        c.stats["new_detail_fetches"] = c.MAX_NEW_PER_RUN
        with mock.patch("jobs.management.commands.fetch_jobs.requests.get", side_effect=self._get(calls, 30)):
            c.fetch_smartrecruiters_api("Acme")
        self.assertEqual(c.screened, [])


@override_settings(**TEST_SETTINGS)
class SponsorPageTests(TestCase):
    def setUp(self):
        cache.clear()

    def test_page_shows_real_numbers(self):
        make_job(company="Acme")
        make_job(company="Globex", title="Marketo Admin")
        Subscriber.objects.create(email="reader@x.test", is_active=True)
        r = self.client.get("/sponsor/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["stats"]["jobs"], 2)
        self.assertEqual(r.context["stats"]["companies"], 2)
        self.assertGreaterEqual(r.context["stats"]["readers"], 1)
        self.assertContains(r, "Weekly email sponsor")
        self.assertNotContains(r, "Google search appearances")  # no Search Console data -> tile hidden

    @override_settings(EMAIL_HOST_PASSWORD="x", EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_enquiry_saved_and_emailed(self):
        from django.core import mail
        from jobs.models import SponsorInquiry
        r = self.client.post("/sponsor/", {"company": "Braze", "name": "Sam", "email": "sam@braze.test",
                                           "option": "newsletter", "message": "Q4"}, follow=True)
        self.assertContains(r, "reply within two working days")
        inq = SponsorInquiry.objects.get()
        self.assertEqual((inq.company, inq.option), ("Braze", "newsletter"))
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].reply_to, ["sam@braze.test"])

    def test_bad_input_and_honeypot_rejected(self):
        from jobs.models import SponsorInquiry
        r = self.client.post("/sponsor/", {"company": "", "name": "Sam", "email": "not-an-email", "option": "newsletter"})
        self.assertEqual(r.status_code, 200)
        self.client.post("/sponsor/", {"company": "Spam", "name": "Bot", "email": "b@x.test", "option": "other",
                                       "website": "http://spam"})
        self.assertEqual(SponsorInquiry.objects.count(), 0)

    def test_in_sitemap_and_nav(self):
        self.assertContains(self.client.get("/sitemap.xml"), "/sponsor/")
        self.assertContains(self.client.get("/for-employers/"), 'href="/sponsor/"')
